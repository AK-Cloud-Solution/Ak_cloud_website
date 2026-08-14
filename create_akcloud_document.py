from datetime import date
from pathlib import Path
import sys

sys.path.insert(0, "/private/tmp/akcloud-docx-deps")

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("AK_Cloud_Enterprise_Technical_Guide.docx")
NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(92, 103, 112)
LIGHT = "E8EEF5"
PALE = "F4F6F9"
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")


def set_font(run, size=None, bold=None, color=None, name="Calibri", italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, 9, color=GRAY)


def add_para(doc, text="", bold_lead=None, style=None, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, 11, True, NAVY)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, 11)
    else:
        r = p.add_run(text)
        set_font(r, 11)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        set_font(p.add_run(item), 11)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.167
        set_font(p.add_run(item), 11)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label + " "), 11, True, NAVY)
    set_font(p.add_run(text), 11)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(header), 10, True, NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), 9.5)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(0.85)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in (
    ("Title", 28, NAVY, 0, 8),
    ("Subtitle", 14, GRAY, 0, 16),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = name != "Subtitle"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(hp.add_run("AK CLOUD ENTERPRISE  |  TECHNICAL GUIDE"), 9, True, GRAY)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(fp.add_run("Internal project documentation  •  "), 9, color=GRAY)
add_field(fp, "PAGE")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(48)
p.paragraph_format.space_after = Pt(8)
r = p.add_run("AK Cloud Enterprise")
set_font(r, 28, True, NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
set_font(p.add_run("Technical Guide, Runbook & Build Verification"), 15, color=GRAY)

meta = doc.add_table(rows=4, cols=2)
meta.style = "Table Grid"
set_table_widths(meta, [1.45, 5.05])
for row, (label, value) in zip(meta.rows, [
    ("Project", "AK Cloud role-based DevOps workspace"),
    ("Stack", "React 18 + TypeScript + Vite 6; Express 5 + MongoDB"),
    ("Verified", str(date.today().strftime("%d %B %Y"))),
    ("Build status", "PASS — production frontend bundle generated"),
]):
    set_cell_shading(row.cells[0], LIGHT)
    set_font(row.cells[0].paragraphs[0].add_run(label), 10, True, NAVY)
    set_font(row.cells[1].paragraphs[0].add_run(value), 10)

doc.add_heading("Executive summary", level=1)
add_para(doc, "AK Cloud Enterprise is a full-stack web application for role-aware access to DevOps tools and a concept workspace. The public experience presents company information and contact content; authenticated users gain access to tool views and role-specific dashboards. A separate workspace concept supports projects, file uploads, meeting records, and mock meeting-minute generation.")
add_callout(doc, "Verification result:", "The frontend production build completed successfully with Vite 6.3.5. It transformed 1,738 modules and emitted the optimized HTML, CSS, and JavaScript assets into build/.")

doc.add_heading("What was verified", level=2)
add_bullets(doc, [
    "Locked frontend dependencies installed with npm ci.",
    "npm run build completed successfully in approximately one second.",
    "Generated bundle: build/index.html, an 83.55 kB CSS asset, and a 468.35 kB JavaScript asset.",
    "Source routes, authentication flow, API client, backend routes, and MongoDB models were inspected.",
])

doc.add_page_break()
doc.add_heading("1. System overview", level=1)
add_table(doc, ["Layer", "Technology", "Responsibility"], [
    ("Presentation", "React 18, TypeScript, Tailwind CSS, Radix UI", "Landing page, dialogs, navigation, dashboards, and workspace screens"),
    ("Build", "Vite 6", "Development server and optimized production bundle"),
    ("Client state", "React context + browser storage", "Authenticated user state and persisted JWT"),
    ("API", "Express 5", "Authentication, files, meetings, and projects"),
    ("Data", "MongoDB + Mongoose", "Users, projects, files, and meeting records"),
    ("Security", "bcryptjs + JWT", "Password hashing and one-day bearer tokens"),
    ("Messaging", "Nodemailer", "Password-reset email through configured Gmail credentials"),
], [1.15, 1.75, 3.6])

doc.add_heading("Application flow", level=2)
add_steps(doc, [
    "A visitor opens the public landing page and can view the home, About Us, and Contact sections.",
    "Tools remain behind a restricted-access prompt until the visitor logs in or registers.",
    "Registration creates a standard user in MongoDB, hashes the password, and returns a JWT.",
    "Login returns the user profile and JWT; the token is stored in localStorage when Remember Me is selected, otherwise in sessionStorage.",
    "The dashboard shown by the frontend depends on the user role: admin or user.",
    "The Concept Demo calls authenticated APIs for projects, file storage, meetings, and generated meeting notes.",
])

doc.add_heading("Role behavior", level=2)
add_table(doc, ["Role", "Observed access"], [
    ("Visitor", "Public landing content; Tools replaced by a restricted-access prompt; no dashboard"),
    ("User", "Authenticated Tools plus User Dashboard; Zoom is presented as available and other app tiles are locked"),
    ("Admin", "Authenticated Tools plus Admin Dashboard; all seven application tiles are presented as available"),
], [1.15, 5.35])

doc.add_page_break()
doc.add_heading("2. Runbook", level=1)
doc.add_heading("Prerequisites", level=2)
add_bullets(doc, [
    "Node.js and npm",
    "MongoDB running locally or an accessible MongoDB Atlas cluster",
    "Optional Gmail application credentials for password-reset email",
])

doc.add_heading("Environment configuration", level=2)
add_para(doc, "Create server/.env with values equivalent to the following. Replace all example secrets before any shared or production deployment.")
add_table(doc, ["Variable", "Example / purpose"], [
    ("PORT", "5000"),
    ("MONGO_URI", "mongodb://localhost:27017/akcloud"),
    ("JWT_SECRET", "A long, random secret; required for secure JWT signing"),
    ("EMAIL_USER", "Sender account used by Nodemailer"),
    ("EMAIL_PASS", "Gmail app password or provider credential"),
    ("VITE_API_URL", "Optional frontend override; defaults to http://localhost:5000/api"),
], [1.7, 4.8])

doc.add_heading("Install and start", level=2)
add_steps(doc, [
    "From the repository root, run npm ci.",
    "From server/, run npm ci to install backend dependencies.",
    "Start MongoDB or confirm MONGO_URI can be reached.",
    "Optionally run node server/createAdmin.js once to create the documented admin account.",
    "From the repository root, run npm run dev. This starts Vite and the backend development server concurrently.",
    "Open the Vite address shown in the terminal, normally http://localhost:5173.",
])

doc.add_heading("Production frontend build", level=2)
add_para(doc, "Run npm run build. The generated static frontend is written to build/. Deploy that directory behind a web server and point VITE_API_URL at the deployed API before building.")

doc.add_heading("Build verification record", level=2)
add_table(doc, ["Check", "Result"], [
    ("Command", "npm run build"),
    ("Vite", "6.3.5"),
    ("Modules", "1,738 transformed"),
    ("Output", "build/index.html and hashed CSS/JS assets"),
    ("Status", "PASS"),
    ("Warnings", "Browser compatibility datasets are stale; update during routine maintenance"),
], [1.6, 4.9])

doc.add_page_break()
doc.add_heading("3. API reference", level=1)
add_para(doc, "Protected routes expect the JWT in the auth-token request header. Records are scoped to req.user.id by the route handlers.")
add_table(doc, ["Method", "Endpoint", "Purpose", "Auth"], [
    ("POST", "/api/auth/register", "Create a user and return profile + JWT", "No"),
    ("POST", "/api/auth/login", "Validate credentials and return profile + JWT", "No"),
    ("GET", "/api/auth/me", "Return the current user", "Yes"),
    ("POST", "/api/auth/forgot-password", "Create token and send reset email", "No"),
    ("POST", "/api/auth/reset-password/:token", "Replace password for a valid reset token", "No"),
    ("POST", "/api/projects", "Create a project", "Yes"),
    ("GET", "/api/projects", "List the caller’s projects", "Yes"),
    ("DELETE", "/api/projects/:id", "Delete the caller’s project", "Yes"),
    ("POST", "/api/files/upload", "Upload one file, maximum 10 MB", "Yes"),
    ("GET", "/api/files", "List files; optional projectId filter", "Yes"),
    ("DELETE", "/api/files/:id", "Delete file metadata and local file", "Yes"),
    ("POST", "/api/meetings/generate", "Create a mock summary/action list", "Yes"),
    ("POST", "/api/meetings", "Save meeting content", "Yes"),
    ("GET", "/api/meetings", "List the caller’s meetings", "Yes"),
    ("DELETE", "/api/meetings/:id", "Delete the caller’s meeting", "Yes"),
], [0.65, 2.05, 3.1, 0.7])

doc.add_heading("Data entities", level=2)
add_bullets(doc, [
    "User — username, email, password hash, role, and password-reset fields.",
    "Project — name, optional description, owner, and timestamps.",
    "File — generated filename, original name, MIME type, size, local path, owner, and optional project.",
    "Meeting — title, transcript, summary, action items, owner, and timestamps.",
])

doc.add_page_break()
doc.add_heading("4. Operational notes and risks", level=1)
add_callout(doc, "Important:", "The frontend build is verified, but a complete backend runtime requires MongoDB and server environment variables. These external services were not provisioned as part of this document run.")

doc.add_heading("Security and reliability findings", level=2)
add_table(doc, ["Priority", "Finding", "Recommended action"], [
    ("Critical", "JWT signing falls back to the literal secretkey when JWT_SECRET is absent.", "Fail startup when JWT_SECRET is missing; rotate any exposed tokens."),
    ("High", "npm reported 7 dependency vulnerabilities: 5 high and 2 critical.", "Run npm audit, assess direct/transitive packages, and upgrade with regression testing."),
    ("High", "CORS accepts all origins.", "Restrict origins, methods, and headers per deployed environment."),
    ("High", "Uploaded files are stored locally with original names appended and no file-type allowlist.", "Validate MIME/signature, scan content, use object storage, and enforce quotas."),
    ("Medium", "Forgot/reset password URLs are hard-coded to localhost in UI/API code.", "Use environment-driven frontend and API base URLs."),
    ("Medium", "The /me handler does not explicitly handle a deleted user before dereferencing.", "Return 404/401 when the token’s user no longer exists."),
    ("Medium", "Meeting “AI” generation is keyword-based mock logic.", "Label it as mock functionality or connect a reviewed summarization service."),
    ("Low", "No automated test script is defined for the backend.", "Add unit and integration tests for auth, ownership checks, uploads, and reset tokens."),
], [0.7, 2.55, 3.25])

doc.add_heading("Documentation differences found", level=2)
add_bullets(doc, [
    "The existing USER_GUIDE says authentication is frontend-only and simulated; the current code calls a real backend and MongoDB.",
    "The existing guide says visitors cannot view About Us; the current App.tsx makes About Us public.",
    "The admin seed script creates admin/admin123, but no equivalent seed script for user/user123 was found.",
    "The Concept Demo includes projects, file management, settings, and meeting-minute features that are not covered by the older guide.",
])

doc.add_heading("Recommended acceptance checks", level=2)
add_steps(doc, [
    "Register a new user and verify duplicate username/email handling.",
    "Test Remember Me and session-only login across browser restarts.",
    "Verify visitor, user, and admin UI access paths.",
    "Create and delete a project; upload, filter, and delete a file.",
    "Generate, save, list, and delete meeting notes.",
    "Request a password reset and verify expiry after one hour.",
    "Confirm one user cannot read or delete another user’s records.",
])

doc.core_properties.title = "AK Cloud Enterprise Technical Guide"
doc.core_properties.subject = "Technical documentation and verified build record"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.keywords = "AK Cloud, React, Express, MongoDB, runbook, API"
doc.save(OUT)
print(OUT.resolve())
